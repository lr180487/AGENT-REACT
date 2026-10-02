from __future__ import annotations
import json, uuid
from pathlib import Path
from typing import Any
from fastapi import APIRouter, File, HTTPException, UploadFile, status, Depends
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from app.config import settings
from app.db.database import get_db
from app.db.models import Document, DocumentChunk
from app.db.repositories.document_repository import DocumentRepository
from app.rag.embeddings import EmbeddingService

router = APIRouter(prefix="/documents", tags=["Documents"])
ALLOWED={".pdf",".csv",".docx",".txt",".xlsx",".md",".json"}
UPLOAD_DIR=Path(settings.UPLOAD_DIR); UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

class DocumentResponse(BaseModel):
    id:str; filename:str; original_filename:str; extension:str; content_type:str|None=None; size:int; status:str; path:str; uploaded_at:str|None=None; chunks:int=0; embedded:bool=False; indexed:bool=False
class DocumentProcessRequest(BaseModel): document_id:str=Field(...); create_chunks:bool=True; create_embeddings:bool=True; index_vectorstore:bool=True
class DocumentProcessResponse(BaseModel):
    success:bool; document_id:str; status:str; text_length:int=0; chunks:int=0; embedded:bool=False; indexed:bool=False; message:str; metadata:dict[str,Any]={}; timestamp:str

def to_response(d): return DocumentResponse(id=d.id, filename=d.filename, original_filename=d.original_filename, extension=d.extension, content_type=d.content_type, size=d.size, status=d.status, path=d.path, uploaded_at=d.created_at.isoformat() if d.created_at else None, chunks=len(d.chunks), embedded=any(c.embedding_json for c in d.chunks), indexed=any(c.embedding_vector is not None for c in d.chunks))

def extract(path: Path, ext: str)->str:
    if ext in {".txt",".md"}: return path.read_text(encoding="utf-8", errors="ignore")
    if ext==".json": return json.dumps(json.loads(path.read_text(encoding="utf-8")), ensure_ascii=False, indent=2)
    if ext==".csv":
        import csv
        with path.open(encoding="utf-8", errors="ignore", newline="") as f: return "\n".join(" | ".join(row) for row in csv.reader(f))
    if ext==".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if ext==".docx":
        from docx import Document as DocxDocument
        return "\n".join(p.text for p in DocxDocument(str(path)).paragraphs)
    if ext==".xlsx":
        from openpyxl import load_workbook
        wb=load_workbook(path, read_only=True, data_only=True); out=[]
        for ws in wb.worksheets:
            out.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True): out.append(" | ".join("" if x is None else str(x) for x in row))
        return "\n".join(out)
    raise ValueError(f"Loader no soportado: {ext}")

def chunks(text, size=None, overlap=None):
    size=size or settings.CHUNK_SIZE; overlap=overlap if overlap is not None else settings.CHUNK_OVERLAP
    out=[]; start=0
    while start<len(text):
        end=min(start+size,len(text)); c=text[start:end].strip()
        if c: out.append(c)
        if end>=len(text): break
        start=end-overlap
    return out

@router.post("/upload", response_model=DocumentResponse, status_code=201)
async def upload(file: UploadFile=File(...), db: Session=Depends(get_db)):
    name=Path(file.filename or "document").name; ext=Path(name).suffix.lower()
    if ext not in ALLOWED: raise HTTPException(400, f"Extensión no permitida: {ext}")
    doc_id=str(uuid.uuid4()); dest=UPLOAD_DIR/f"{doc_id}{ext}"; total=0
    with dest.open("wb") as f:
        while data:=await file.read(1024*1024):
            total+=len(data)
            if total>settings.MAX_FILE_SIZE_MB*1024*1024: dest.unlink(missing_ok=True); raise HTTPException(413,"Archivo demasiado grande")
            f.write(data)
    d=Document(id=doc_id, filename=dest.name, original_filename=name, extension=ext, content_type=file.content_type, path=str(dest), size=total, status="uploaded")
    DocumentRepository(db).add(d); return to_response(d)

@router.get("", response_model=list[DocumentResponse])
def list_documents(limit:int=50, db:Session=Depends(get_db)): return [to_response(x) for x in DocumentRepository(db).list(limit)]
@router.get("/{document_id}", response_model=DocumentResponse)
def get_document(document_id:str, db:Session=Depends(get_db)):
    d=DocumentRepository(db).get(document_id)
    if not d: raise HTTPException(404,"Documento no encontrado")
    return to_response(d)

@router.post("/{document_id}/process", response_model=DocumentProcessResponse)
async def process_document(document_id:str, request:DocumentProcessRequest, db:Session=Depends(get_db)):
    repo=DocumentRepository(db); d=repo.get(document_id)
    if not d: raise HTTPException(404,"Documento no encontrado")
    text=extract(Path(d.path),d.extension); cs=chunks(text) if request.create_chunks else []
    vectors=await EmbeddingService().embed(cs) if request.create_embeddings and cs else []
    rows=[]
    for i,c in enumerate(cs): rows.append(DocumentChunk(id=str(uuid.uuid4()), document_id=d.id, chunk_index=i, content=c, embedding_json=vectors[i] if vectors else None, embedding_vector=vectors[i] if vectors and db.bind.dialect.name=="postgresql" else None))
    repo.replace_chunks(d.id, rows); d.status="indexed" if rows else "processed"; db.commit(); db.refresh(d)
    return DocumentProcessResponse(success=True,document_id=d.id,status=d.status,text_length=len(text),chunks=len(rows),embedded=bool(vectors),indexed=bool(vectors and db.bind.dialect.name=="postgresql"),message="Documento procesado e indexado",metadata={"loader":d.extension},timestamp=__import__('datetime').datetime.now(__import__('datetime').timezone.utc).isoformat())

@router.get("/{document_id}/download")
def download(document_id:str, db:Session=Depends(get_db)):
    d=DocumentRepository(db).get(document_id)
    if not d or not Path(d.path).exists(): raise HTTPException(404,"Documento no encontrado")
    return FileResponse(d.path,filename=d.original_filename,media_type=d.content_type or "application/octet-stream")

@router.delete("/{document_id}")
def delete(document_id:str, db:Session=Depends(get_db)):
    d=DocumentRepository(db).get(document_id)
    if not d: raise HTTPException(404,"Documento no encontrado")
    Path(d.path).unlink(missing_ok=True); db.delete(d); db.commit(); return {"success":True,"document_id":document_id}
